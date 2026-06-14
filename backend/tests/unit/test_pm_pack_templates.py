"""Phase 104 Plan 01 Task 3 — template-shape oracle test for the PM pack templates.

Asserts the two seeded ``.docx`` templates under ``scripts/pm-pack/templates/`` carry the
exact placeholder contract the harness emit path fills, and that the Risk Register's
inline P×I Score expression renders a non-blank score from WORDED inputs with ZERO
backend code.

Fixture-shape fidelity (CONFIRMED against build_context, ``tool_dispatcher.py:1225-1240``):
the production render context is built of plain **dicts**, NOT objects — ``build_context``
maps each collection row through ``_build_row`` -> ``{col: _cell(cited)}`` where
``_cell(...)`` returns a dict ``{"value": ..., "source_chunk_id": ..., "source_doc": ...,
"source_page": ...}`` (``tool_dispatcher.py:1207-1222``). In Jinja, ``r.probability.value``
resolves to ``row["probability"]["value"]`` (attribute access falls back to dict item
access). So the inline Score expression and this test's ``_cell(v)`` dicts are the SAME
shape build_context produces — there is NO SimpleNamespace/object access in this path, so
a passing render here proves the LIVE render path, not a divergent fixture.

Tests 3+4 render through ``docxtpl.DocxTemplate`` + a
``jinja2.sandbox.SandboxedEnvironment(autoescape=True)`` — the exact driver the production
``render_docx_template`` uses (``tool_dispatcher.py:1261-1273``).
"""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("docx")
pytest.importorskip("docxtpl")

# Templates live at the repo root under scripts/pm-pack/templates/.
# This file is backend/tests/unit/test_pm_pack_templates.py -> parents[3] == repo root.
REPO_ROOT = Path(__file__).resolve().parents[3]
TEMPLATES_DIR = REPO_ROOT / "scripts" / "pm-pack" / "templates"
STATUS_TEMPLATE = TEMPLATES_DIR / "weekly-status-report.docx"
RISK_TEMPLATE = TEMPLATES_DIR / "risk-register.docx"

# Wave-0 ordering safety: if the templates were not built yet, skip rather than fail.
_templates_built = STATUS_TEMPLATE.exists() and RISK_TEMPLATE.exists()
skip_if_unbuilt = pytest.mark.skipif(
    not _templates_built,
    reason="PM pack templates not built (run scripts/pm-pack/make_pm_templates.py)",
)

STATUS_SCALARS = {
    "project_name",
    "reporting_period",
    "overall_rag_status",
    "summary",
    "accomplishments",
    "planned_next",
    "risks_blockers",
    "milestones",
}

RISK_CITED_COLUMNS = [
    "id",
    "description",
    "category",
    "probability",
    "impact",
    "owner",
    "mitigation",
    "status",
]


def _docx_full_text(path: Path) -> str:
    """All paragraph + table-cell text of a docx, concatenated."""
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _cell(v):
    """A VERBATIM mirror of the production ``_cell`` dict shape
    (tool_dispatcher.py build_context, :1207-1222) — a dict with a ``value`` key, NOT a
    SimpleNamespace. This is exactly what build_context produces for each cited cell, so a
    render that passes here proves the live render path."""
    return {"value": v, "source_chunk_id": None, "source_doc": None, "source_page": None}


_SCORE_COL_INDEX = 5  # RISK_HEADERS: id, desc, category, probability, impact, [Score], ...


def _render_risk_score_cell(row: dict, tmp_path: Path) -> str:
    """Render the risk template with a single row through the production driver shape
    (docxtpl + SandboxedEnvironment) and return the rendered Score cell text exactly."""
    from docx import Document
    from docxtpl import DocxTemplate
    from jinja2.sandbox import SandboxedEnvironment

    # All scalars blanked (the cell under test is the Score cell of the single row).
    context: dict = {key: _cell("") for key in STATUS_SCALARS}
    context["project_name"] = _cell("")
    context["reporting_period"] = _cell("")
    context["rows"] = [row]

    doc = DocxTemplate(str(RISK_TEMPLATE))
    jenv = SandboxedEnvironment(autoescape=True)  # mirrors render_docx_template
    doc.render(context, jinja_env=jenv)

    out = tmp_path / "rendered-risk-register.docx"
    doc.save(str(out))
    rendered = Document(str(out))
    # The single grown content row is row index 1 (row 0 is the header).
    content_row = rendered.tables[0].rows[1]
    return content_row.cells[_SCORE_COL_INDEX].text.strip()


@skip_if_unbuilt
def test_status_template_exposes_named_scalar_keys():
    """Test 1: parse_docx_template_variables over the status template returns a key set
    that is a superset of the 8 named status scalar keys (the EmitFieldMap scalar contract)."""
    from app.services.template_render_service import parse_docx_template_variables

    oracle = parse_docx_template_variables(STATUS_TEMPLATE.read_bytes())
    assert oracle is not None, "status template carried no parseable tokens"
    got = set(oracle["scalars"])
    missing = STATUS_SCALARS - got
    assert not missing, f"status template missing scalar keys: {sorted(missing)}"


@skip_if_unbuilt
def test_risk_template_has_loop_and_cited_columns_but_no_bare_score():
    """Test 2: the risk template grows rows via {%tr for r in rows %} and exposes the 8
    cited {{ r.<col>.value }} columns; the bare {{ r.score }} token is ABSENT (the Score is
    an inline P×I expression, not a model-emitted / driver-computed field)."""
    text = _docx_full_text(RISK_TEMPLATE)

    assert "{%tr for r in rows %}" in text, "risk template missing the {%tr for%} row-growth loop"
    for col in RISK_CITED_COLUMNS:
        token = "{{ r.%s.value }}" % col
        assert token in text, f"risk template missing cited column token {token}"

    assert "{{ r.score }}" not in text, (
        "risk template must NOT use the bare {{ r.score }} token (renders blank in prod — "
        "build_context applies no numeric_hook); the Score must be an inline P×I expression"
    )
    # The inline Score expression dereferences the SAME cited cells the model emits.
    assert "r.probability.value" in text
    assert "r.impact.value" in text


@skip_if_unbuilt
def test_inline_score_renders_six_for_high_times_medium(tmp_path):
    """Test 3: a worded High x Medium renders Score "6" (High=3 x Med=2) via the inline
    Jinja P×I expression — proving D-104-4's worded->numeric score mapping with NO backend
    code. The fixture row dicts are the verbatim build_context _cell shape, so this proves
    the live render path."""
    row = {
        "id": _cell("M-01"),
        "description": _cell("Duplicate customer records"),
        "category": _cell("Data"),
        "probability": _cell("High"),
        "impact": _cell("Medium"),
        "owner": _cell("Priya Nair"),
        "mitigation": _cell("Dedup pass + validation gate"),
        "status": _cell("In progress"),
    }
    score = _render_risk_score_cell(row, tmp_path)
    assert score == "6", f"expected inline P×I Score '6' (High=3 x Med=2); got {score!r}"


@skip_if_unbuilt
def test_inline_score_degrades_to_zero_for_unmapped_value(tmp_path):
    """Test 4 (negative control / honest degrade): an unmapped probability ("Unknown")
    renders Score "0" via the .get default and never crashes the render."""
    row = {
        "id": _cell("M-99"),
        "description": _cell("Unmapped probability"),
        "category": _cell("Other"),
        "probability": _cell("Unknown"),
        "impact": _cell("High"),
        "owner": _cell("N/A"),
        "mitigation": _cell("N/A"),
        "status": _cell("Open"),
    }
    score = _render_risk_score_cell(row, tmp_path)
    assert score == "0", f"expected inline Score '0' for unmapped probability; got {score!r}"
