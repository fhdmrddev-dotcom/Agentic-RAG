"""Phase 291 (AUTH-02, plan 291-01 Task 2) — the Word fillable COPY (D-26, D-27, D-28).

The AI names spots (an anchor: the label text and the kind of blank). The SERVER decides
where each one is, writes ``{{ key.value }}`` into a COPY, and then re-reads the copy to
prove every ticked key appears exactly once. A label it cannot place exactly is said in
words and left blank, never guessed. The original bytes are never modified.

Fixtures are python-docx documents built in the test; the proposer is a fake.
"""

from __future__ import annotations

import hashlib
import io
from types import SimpleNamespace

import pytest

_SETTINGS = SimpleNamespace(multimodal_max_vision_calls=100, vision_max_pages=50)


def _blanks_docx(signature_times: int = 2, extra: list[str] | None = None) -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Vendor onboarding form")
    p = d.add_paragraph()
    r = p.add_run("Name: ")
    r.bold = True
    p.add_run("_____")
    p.add_run("_____")
    d.add_paragraph("Date of birth: ________")
    d.add_paragraph("Keep this line exactly as it is.")
    for _ in range(signature_times):
        d.add_paragraph("Signature: __________")
    d.add_paragraph("Phone")
    d.add_paragraph("__________")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Company"
    t.rows[0].cells[1].text = ""
    for line in extra or []:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _paragraph_texts(data: bytes) -> list[str]:
    import docx

    return [p.text for p in docx.Document(io.BytesIO(data)).paragraphs]


def _anchor(key, kind, label):
    return {"key": key, "anchor_kind": kind, "anchor_label": label}


COUPLDNT_TWICE = (
    "Couldn't place “Signature”: that label appears twice in the document. "
    "It stays blank in the filled copy."
)
COUPLDNT_THREE = (
    "Couldn't place “Signature”: that label appears 3 times in the document. "
    "It stays blank in the filled copy."
)
SPLIT_PHONE = (
    "Couldn't place “Phone”: its blank is split in a way that can't be found "
    "exactly. It stays blank in the filled copy."
)
COPY_REFUSED = (
    "This document already contains {{ }} or {% %} marks that aren't fields, so a fillable "
    "copy can't be made from it. Remove them and attach it again."
)


def test_line_blank_gets_value_token_and_keeps_run_formatting():
    import docx

    from app.services.describe_files import make_fillable_copy

    original = _blanks_docx()
    res = make_fillable_copy(original, [_anchor("name", "line", "Name")])
    assert res.refused is None
    assert res.placed == ["name"]
    d = docx.Document(io.BytesIO(res.copy))
    para = d.paragraphs[1]
    assert para.text == "Name: {{ name.value }}"
    assert para.runs[0].text == "Name: " and para.runs[0].bold is True
    before = _paragraph_texts(original)
    after = _paragraph_texts(res.copy)
    assert len(before) == len(after)
    for i, (a, b) in enumerate(zip(before, after)):
        if i != 1:
            assert a == b


def test_empty_cell_right_of_label_gets_the_token():
    import docx

    from app.services.describe_files import make_fillable_copy

    res = make_fillable_copy(_blanks_docx(), [_anchor("company", "cell", "Company")])
    assert res.placed == ["company"]
    d = docx.Document(io.BytesIO(res.copy))
    assert d.tables[0].rows[0].cells[1].text == "{{ company.value }}"
    assert d.tables[0].rows[0].cells[0].text == "Company"


def test_duplicate_label_is_not_placed_and_says_twice():
    from app.services.describe_files import make_fillable_copy

    res = make_fillable_copy(_blanks_docx(2), [_anchor("signature", "line", "Signature")])
    assert res.placed == []
    assert res.unplaced == [("Signature", COUPLDNT_TWICE)]


def test_three_appearances_say_three_times():
    from app.services.describe_files import make_fillable_copy

    res = make_fillable_copy(_blanks_docx(3), [_anchor("signature", "line", "Signature")])
    assert res.unplaced == [("Signature", COUPLDNT_THREE)]


def test_split_blank_says_it_cannot_be_found_exactly():
    from app.services.describe_files import make_fillable_copy

    res = make_fillable_copy(_blanks_docx(), [_anchor("phone", "line", "Phone")])
    assert res.placed == []
    assert res.unplaced == [("Phone", SPLIT_PHONE)]


def test_every_ticked_key_appears_exactly_once_and_parser_sees_it():
    from app.services.describe_files import count_token_occurrences, make_fillable_copy
    from app.services.template_render_service import parse_docx_template_variables

    anchors = [
        _anchor("name", "line", "Name"),
        _anchor("date_of_birth", "line", "Date of birth"),
        _anchor("company", "cell", "Company"),
    ]
    res = make_fillable_copy(_blanks_docx(), anchors)
    assert sorted(res.placed) == ["company", "date_of_birth", "name"]
    parsed = parse_docx_template_variables(res.copy)
    assert set(parsed["scalars"]) == {"company", "date_of_birth", "name"}
    for key in res.placed:
        assert count_token_occurrences(res.copy, key) == 1


def test_a_key_anchored_twice_is_refused_for_the_second_occurrence():
    from app.services.describe_files import count_token_occurrences, make_fillable_copy

    res = make_fillable_copy(
        _blanks_docx(),
        [_anchor("name", "line", "Name"), _anchor("name", "line", "Date of birth")],
    )
    assert res.placed == ["name"]
    assert count_token_occurrences(res.copy, "name") == 1
    assert [label for label, _ in res.unplaced] == ["Date of birth"]
    assert "Date of birth: ________" in _paragraph_texts(res.copy)


def test_unticked_blanks_get_no_token_and_keep_their_underscores():
    from app.services.describe_files import make_fillable_copy

    res = make_fillable_copy(_blanks_docx(), [_anchor("name", "line", "Name")])
    texts = _paragraph_texts(res.copy)
    assert "Date of birth: ________" in texts
    assert "date_of_birth" not in "".join(texts)


def test_invalid_and_reserved_keys_are_dropped_in_words():
    from app.services.describe_files import make_fillable_copy

    res = make_fillable_copy(
        _blanks_docx(),
        [_anchor("9bad", "line", "Name"), _anchor("kickoff_prompt", "line", "Date of birth")],
    )
    assert res.placed == []
    assert [label for label, _ in res.unplaced] == ["Name", "Date of birth"]
    assert all(s.startswith("Couldn't place") for _, s in res.unplaced)


@pytest.mark.parametrize("stray", ["Notes {{ not_ours }}", "{% if x %}maybe{% endif %}"])
def test_original_with_template_marks_is_refused_in_words(stray):
    from app.services.describe_files import make_fillable_copy

    res = make_fillable_copy(_blanks_docx(extra=[stray]), [_anchor("name", "line", "Name")])
    assert res.copy is None
    assert res.refused == COPY_REFUSED
    assert res.placed == []


def test_original_bytes_are_never_modified():
    from app.services.describe_files import make_fillable_copy

    original = _blanks_docx()
    before = hashlib.md5(original).hexdigest()
    make_fillable_copy(original, [_anchor("name", "line", "Name"), _anchor("company", "cell", "Company")])
    assert hashlib.md5(original).hexdigest() == before


def test_copy_renders_values_not_a_dict_repr_and_is_residual_clean(tmp_path):
    from app.services.describe_files import make_fillable_copy
    from app.services.template_render_service import render_docx_template, residual_tags_in

    res = make_fillable_copy(
        _blanks_docx(), [_anchor("name", "line", "Name"), _anchor("company", "cell", "Company")]
    )
    src = tmp_path / "copy.docx"
    src.write_bytes(res.copy)
    out = tmp_path / "out.docx"

    def _cited(v):
        return {"value": v, "source_chunk_id": "c1", "source_doc": "d", "source_page": 1}

    rendered = render_docx_template(
        str(src), {"name": _cited("Ada Lovelace"), "company": _cited("Acme & Co")}, str(out)
    )
    assert rendered["rendered"] is True
    assert residual_tags_in(str(out), "docx") == []
    import docx

    d = docx.Document(str(out))
    body = "\n".join(p.text for p in d.paragraphs)
    assert "Name: Ada Lovelace" in body
    assert "'value'" not in body and "{'" not in body
    assert d.tables[0].rows[0].cells[1].text == "Acme & Co"


class _FakeProposer:
    def __init__(self, fields):
        self.fields = fields
        self.calls = []

    def __call__(self, text, candidates):
        self.calls.append((text, list(candidates)))
        return list(self.fields)


def test_read_blanks_arm_shows_each_line_and_the_unplaced_duplicate():
    from app.services import describe_files as df

    proposer = _FakeProposer(
        [
            {"key": "name", "label": "Name", "type": "text", "anchor_kind": "line", "anchor_label": "Name"},
            {
                "key": "date_of_birth",
                "label": "Date of birth",
                "type": "text",
                "anchor_kind": "line",
                "anchor_label": "Date of birth",
            },
            {"key": "signature", "label": "Signature", "type": "text", "anchor_kind": "line", "anchor_label": "Signature"},
            {"key": "company", "label": "Company", "type": "text", "anchor_kind": "cell", "anchor_label": "Company"},
        ]
    )
    original = _blanks_docx()
    res = df.read_describe_file(
        original, ".docx", "vendor-form.docx", "form", app_settings=_SETTINGS, propose=proposer
    )
    assert res.arm == "blanks"
    by = {f["key"]: f for f in res.fields}
    assert set(by) == {"name", "date_of_birth", "company"}
    assert all(f["origin"] == "ai_proposed" for f in res.fields)
    assert by["name"]["line_before"] == "Name: " and by["name"]["line_after"] == ""
    assert by["name"]["anchor_kind"] == "line" and by["name"]["anchor_label"] == "Name"
    assert by["date_of_birth"]["type"] == "date"  # D-11: an obvious Date label
    assert by["company"]["anchor_kind"] == "cell"
    assert res.unplaced == [{"label": "Signature", "sentence": COUPLDNT_TWICE}]
    # The proposer saw the blank candidates the server found.
    kinds = {(c["anchor_kind"], c["anchor_label"]) for c in proposer.calls[0][1]}
    assert ("line", "Name") in kinds and ("cell", "Company") in kinds
    # The read is a dry run: the original is untouched.
    assert "{{" not in "".join(_paragraph_texts(original))


def test_read_docx_with_stray_marks_and_no_tokens_is_copy_refused():
    from app.services import describe_files as df

    res = df.read_describe_file(
        _blanks_docx(extra=["Total {% weird"]),
        ".docx",
        "f.docx",
        "form",
        app_settings=_SETTINGS,
        propose=_FakeProposer([]),
    )
    assert res.arm == "copy_refused"
    assert res.refused == COPY_REFUSED
