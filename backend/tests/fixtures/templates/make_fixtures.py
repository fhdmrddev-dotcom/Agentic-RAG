"""Deterministic generator for the four Phase 101 OOXML template fixtures.

Phase 101 Plan 01 (Wave 0, TMPL-02/TMPL-03). These are TEST fixtures — small,
synthetic, regenerable. They carry only placeholder tokens ({{...}} / {%tr %})
and synthetic data — no PII, no secrets (threat T-101-01-03). They are loaded
ONLY in the offline TEST tier; they never reach a production input path and are
never executed (T-101-01-02).

Run from the repo root to (re)generate all four binaries in-place:

    backend/venv/Scripts/python.exe backend/tests/fixtures/templates/make_fixtures.py

The four fixtures and what each exercises:

  1. risk-register.docx        — PORT of scripts/spike-097/make_template.py: the
     trusted docxtpl {%tr for r in rows %} variable-row register + body/scalar
     {{ project_name.value }} / {{ report_date.value }} tags. The coverage oracle
     DocxTemplate(...).get_undeclared_template_variables() must be a superset of
     {project_name, report_date, rows} (the trusted-path render-grows-rows probe).

  2. arbitrary-split-token.docx — the Pitfall-1 reproduction: a literal token
     {{client_name}} intentionally FRAGMENTED across three <w:r> runs (the
     run-split silent-miss the arbitrary non-Jinja run-merge engine must
     reassemble), plus a clean whole-in-one-run {{report_title}} control. No
     Jinja {%...%} blocks — the arbitrary scalar-only path.

  3. table-deck.pptx          — one slide with a 2x2 table carrying {{metric}} /
     {{value}} scalar tokens + a {{deck_title}} title. Exercises the python-pptx
     "cannot grow tables" documented limit (Plan 03 verdict note); scalar fill
     must round-trip.

  4. chart-book.xlsx          — one worksheet with scalar token cells (A1={{title}},
     B2 numeric) + a real BarChart (the openpyxl chart-strip-on-save limit Plan 03
     documents) + a merged-cell range A4:B4 anchored {{merged_note}} (the
     merged-cell mis-write failure mode, SC#4 #6).

This generator is idempotent — re-running overwrites the four binaries cleanly.
Do NOT hand-edit the binaries; regenerate from here.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.util import Inches

OUT_DIR = Path(__file__).resolve().parent

# ── 1. risk-register.docx (trusted docxtpl {%tr %} template) ──────────────────
# Port of scripts/spike-097/make_template.py. Header labels + body Jinja tags in
# the SAME column order is the contract.
_HEADERS = [
    "Risk ID", "Cause", "Event", "Effect",
    "P", "I", "Score", "Response", "Owner", "Status",
]
_BODY_TAGS = [
    "{{ r.risk_id.value }}",
    "{{ r.cause.value }}",
    "{{ r.event.value }}",
    "{{ r.effect.value }}",
    "{{ r.probability.value }}",
    "{{ r.impact.value }}",
    "{{ r.score }}",                 # render-time compute — NO .value
    "{{ r.response_strategy.value }}",
    "{{ r.owner.value }}",
    "{{ r.status.value }}",
]
assert len(_HEADERS) == len(_BODY_TAGS), "header/body column count mismatch"


def build_risk_register() -> Path:
    """A docx with a {%tr for r in rows %} repeat row + scalar {{ }} tags.

    docxtpl replaces each {%tr %}-bearing <w:tr> row with the bare {% ... %}
    statement, so the loop is authored as THREE rows (open / content / close) +
    the header row.
    """
    out = OUT_DIR / "risk-register.docx"
    doc = Document()

    doc.add_heading("Risk Register", level=1)
    # Scalar tags — each in its own single-run paragraph.
    doc.add_paragraph("Project: {{ project_name.value }}")
    doc.add_paragraph("Date: {{ report_date.value }}")
    doc.add_paragraph("")  # spacer

    table = doc.add_table(rows=4, cols=len(_HEADERS))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for i, label in enumerate(_HEADERS):
        header_cells[i].text = label

    table.rows[1].cells[0].text = "{%tr for r in rows %}"

    content_cells = table.rows[2].cells
    for i, tag in enumerate(_BODY_TAGS):
        content_cells[i].text = tag

    table.rows[3].cells[0].text = "{%tr endfor %}"

    doc.save(str(out))
    return out


# ── 2. arbitrary-split-token.docx (Pitfall 1 run-fragmentation reproduction) ──
def build_arbitrary_split_token() -> Path:
    """A docx where {{client_name}} is split across three <w:r> runs.

    Each paragraph.add_run() emits a separate <w:r>, so writing the token in
    three fragments ("{{cli" / "ent_" / "name}}") reproduces Word's mid-token
    run split. The concatenated paragraph text contains the whole token, but no
    single run.text does — exactly the silent-miss the arbitrary run-merge
    engine must reassemble. A second clean paragraph holds a whole-in-one-run
    {{report_title}} control. No Jinja {%...%} blocks (scalar-only path).
    """
    out = OUT_DIR / "arbitrary-split-token.docx"
    doc = Document()

    # First paragraph: the fragmented token across 3 runs.
    p1 = doc.add_paragraph()
    p1.add_run("{{cli")
    p1.add_run("ent_")
    p1.add_run("name}}")

    # Second paragraph: the un-fragmented control token in one run.
    p2 = doc.add_paragraph()
    p2.add_run("{{report_title}}")

    doc.save(str(out))
    return out


# ── 3. table-deck.pptx (scalar table fill + python-pptx grow-table limit) ─────
def build_table_deck() -> Path:
    """A pptx with a 2x2 table carrying {{metric}}/{{value}} scalar tokens + a
    {{deck_title}} title placeholder."""
    out = OUT_DIR / "table-deck.pptx"
    prs = Presentation()
    # Blank layout (index 6 in the default template) so we control all shapes.
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title text box carrying the scalar title token.
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(9), Inches(1)
    )
    title_box.text_frame.text = "{{deck_title}}"

    # 2-row x 2-col table with scalar token cells.
    rows, cols = 2, 2
    table_shape = slide.shapes.add_table(
        rows, cols, Inches(0.5), Inches(1.5), Inches(9), Inches(2)
    )
    table = table_shape.table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "{{metric}}"
    table.cell(1, 1).text = "{{value}}"

    prs.save(str(out))
    return out


# ── 4. chart-book.xlsx (scalar fill + openpyxl chart-strip + merged-cell) ─────
def build_chart_book() -> Path:
    """An xlsx with scalar token cells, a real BarChart (openpyxl strips charts
    on load+save — the documented limit), and a merged-cell range A4:B4 anchored
    by a {{merged_note}} token."""
    out = OUT_DIR / "chart-book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Scalar token cell + numeric data the chart references.
    ws["A1"] = "{{title}}"
    ws["A2"] = "Q1"
    ws["B2"] = 10
    ws["A3"] = "Q2"
    ws["B3"] = 20

    # Merged-cell range with the anchor carrying the token (SC#4 #6).
    ws.merge_cells("A4:B4")
    ws["A4"] = "{{merged_note}}"

    # A real BarChart so the strip-on-save limit is reproducible (Plan 03).
    chart = BarChart()
    chart.title = "Quarterly"
    data = Reference(ws, min_col=2, min_row=2, max_row=3)
    chart.add_data(data, titles_from_data=False)
    ws.add_chart(chart, "D2")

    wb.save(str(out))
    return out


def build_all() -> list[Path]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    return [
        build_risk_register(),
        build_arbitrary_split_token(),
        build_table_deck(),
        build_chart_book(),
    ]


if __name__ == "__main__":
    for path in build_all():
        print(f"wrote {path}")
