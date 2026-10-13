"""Skill tools: ``load_skill``, ``save_skill``, ``read_skill_file`` and ``attach_skill_file``, plus
the ONE org-gated skill-visibility resolver (``_resolve_skill_visibility_or``, SEED-125) that
``tools/sandbox.py`` also imports by name.

Moved VERBATIM out of ``tool_dispatcher.py`` in Phase 278 (SEAM-03, a pure move): every binding
here is AST-identical to its PHASE_BASE source, proven by ``tests/unit/test_278_pure_move.py``.
``tool_dispatcher.py`` keeps the registry: it imports the handlers below BY NAME into its static
``_TOOL_REGISTRY`` literal (Extension Contract — never registered dynamically).

⚠ PATCH-WHERE-USED: the bodies in this module resolve these names from THIS module's globals —
``run_in_threadpool``, ``aexec``, ``_resolve_caller_org_ids``, ``build_skill_visibility_or``,
``write_audit_entry``, ``sandbox_manager``, ``lint_description``, ``_get_file_content``,
``get_file_by_path``, ``guess_mime_type``, ``SCRIPT_EXTS``, ``_fetch_owned_document_bytes``. Patch
``app.services.tools.skills.<name>``; a patch of the same name on ``tool_dispatcher`` no longer
reaches the code that calls it (and, because the registry module no longer imports them, raises
``AttributeError`` instead of silently patching nothing).

⚠ ``_fetch_owned_document_bytes`` is defined in ``tools/documents.py`` but CALLED here by the
attach resolver: patch ``app.services.tools.skills._fetch_owned_document_bytes`` for that path.

⛔ Import-cycle rule: no module-level import of ``app.services.tool_dispatcher`` here (the
dispatcher imports this module at load). Cross-family helpers are imported BY NAME from their
``app.services.tools.<family>`` home, never through the dispatcher.
"""
from __future__ import annotations

import io
import json
import logging
import os
from uuid import UUID

from starlette.concurrency import run_in_threadpool

from app.utils.db import aexec
# SEED-125 (CR-01) — the same fail-closed caller-org resolver the SEED-124 folder fix uses,
# reused here to org-gate service-role skill resolution (see _resolve_skill_visibility_or).
from app.utils.folder_utils import _resolve_caller_org_ids
# SEED-125 (CR-01) / Phase 182 (CR-01) — the ONE org-gated skill-visibility predicate, hoisted
# to a shared import-light home so the grounding/canvas seam reuses it instead of copying it.
from app.utils.skill_visibility import build_skill_visibility_or
from app.services.audit_service import write_audit_entry
from app.services.sandbox_service import sandbox_manager
from app.services.skill_lint import lint_description
from app.services.workspace_service import (
    # Phase 151 (FILE-01) — attach_skill_file source #1 (workspace) bytes reader +
    # filename mime guess. get_file_by_path is a workspace_service re-export of the
    # db.workspace fetch (cycle-safe at module level; patch-where-used friendly for tests).
    _get_file_content,
    get_file_by_path,
    guess_mime_type,
)
from app.services.tools.base import ToolContext, ToolResult
from app.services.tools.documents import _fetch_owned_document_bytes
from app.services.tools.runtime_gap import SCRIPT_EXTS

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Phase 151 (FILE-01) — attach_skill_file: save a file onto an OWNED skill from
# one of four sources (workspace / sandbox_output / inline / kb_document, D-05)
# ---------------------------------------------------------------------------
# Inline (source #3) size bound — the weak-model guard. Inline is for small text/config/
# scripts; a model that mangles a large/binary file into a huge `content` string is
# refused here (use the workspace or sandbox_output source for real files). 5 MB is
# generous for text while capping accidental context blow-ups.
_ATTACH_INLINE_MAX_BYTES = 5 * 1024 * 1024


async def _resolve_attach_source_bytes(
    args: dict, ctx: ToolContext, source: str, filename: str
) -> "tuple[bytes, str] | dict":
    """Resolve ``(file_bytes, mime_type)`` from one of the four D-05 sources.

    Returns EITHER ``(bytes, mime)`` on success OR ``{"error": ...}`` (the shared
    honest-failure convention — the caller relays it as a ToolResult error). Every
    blocking Storage/DB/container call is threadpool-wrapped (Pitfall 1).
    """
    import os as _os_local
    import tempfile as _tempfile_local

    # ── #1 workspace: a thread workspace file, read via the shared inline-vs-Storage
    #    reader (`_get_file_content`, already threadpool-safe). ──
    if source == "workspace":
        wpath = (args.get("workspace_path") or "").strip()
        if not wpath:
            return {"error": "source='workspace' requires 'workspace_path'."}
        try:
            file_row = await get_file_by_path(ctx.pool, UUID(ctx.thread_id), wpath)
        except Exception:  # noqa: BLE001 — a bad path/uuid is an honest not-found, never a 500
            file_row = None
        if not file_row:
            return {"error": f"No workspace file at '{wpath}'."}
        try:
            file_bytes = await _get_file_content(ctx.pool, ctx.supabase, file_row)
        except Exception as exc:  # noqa: BLE001
            return {"error": f"Could not read workspace file '{wpath}': {exc}"}
        mime = file_row.get("mime_type") or guess_mime_type(filename)
        return (file_bytes, mime)

    # ── #2 sandbox_output: harvest a named file the agent produced in /sandbox/output,
    #    reusing the harvest idiom (copy the dir out, walk, match basename). ──
    if source == "sandbox_output":
        spath = (args.get("sandbox_path") or "").strip()
        if not spath:
            return {"error": "source='sandbox_output' requires 'sandbox_path'."}
        want = _os_local.path.basename(spath) or spath

        def _harvest() -> "bytes | None":
            session = sandbox_manager.get_or_create(ctx.thread_id)
            try:
                session.execute_command("mkdir -p /sandbox/output")
            except Exception:
                pass  # best-effort; copy_from_runtime 404s if it truly doesn't exist
            with _tempfile_local.TemporaryDirectory() as tmpdir:
                # No trailing slash — Docker's get_archive is strict (sandbox_service.py:264).
                session.copy_from_runtime("/sandbox/output", tmpdir)
                for root, _dirs, files in _os_local.walk(tmpdir):
                    for fn in files:
                        if fn == want:
                            with open(_os_local.path.join(root, fn), "rb") as f:
                                return f.read()
            return None

        try:
            data = await run_in_threadpool(_harvest)
        except Exception as exc:  # noqa: BLE001
            return {"error": f"Could not read sandbox output '{spath}': {exc}"}
        if data is None:
            return {"error": f"No sandbox output file named '{want}' in /sandbox/output."}
        # WR-02 — cap the harvested bytes BEFORE the Storage upload. sandbox_output is the
        # largest, least-trustworthy source (agent-generated bytes) and was the ONLY source
        # with no ceiling; a model can trivially write a huge artifact to /sandbox/output and
        # attach it uncapped (read into backend RAM, then uploaded with no limit). Enforce the
        # shared 5 MB attach ceiling (_ATTACH_INLINE_MAX_BYTES — no new magic number); refuse
        # over-cap with an honest error, no partial upload.
        if len(data) > _ATTACH_INLINE_MAX_BYTES:
            return {"error": (
                f"Sandbox file '{want}' is {len(data) // 1024 // 1024} MB, over the "
                f"{_ATTACH_INLINE_MAX_BYTES // 1024 // 1024} MB attach limit."
            )}
        return (data, guess_mime_type(filename))

    # ── #3 inline: content passed directly as a tool arg (utf-8 text). Weak-model
    #    guards: a non-str / empty / oversized `content` is an honest error, never a
    #    silent empty or context-blowing write. ──
    if source == "inline":
        content = args.get("content")
        if not isinstance(content, str) or content == "":
            return {"error": "source='inline' requires non-empty 'content'."}
        file_bytes = content.encode("utf-8")
        if len(file_bytes) > _ATTACH_INLINE_MAX_BYTES:
            return {"error": (
                f"Inline content is {len(file_bytes) // 1024} KB, over the "
                f"{_ATTACH_INLINE_MAX_BYTES // 1024 // 1024} MB inline limit — attach a "
                "large file via the 'workspace' or 'sandbox_output' source instead."
            )}
        return (file_bytes, guess_mime_type(filename))

    # ── #4 kb_document: REUSE FILE-02's owner-scope resolver (T-03). The doc→skill
    #    data-movement path is owner-scoped on BOTH ends — the resolver's owner→global
    #    gate here, the owner-only skill gate in the caller. Its {"error"} propagates. ──
    if source == "kb_document":
        document_id = (args.get("document_id") or "").strip()
        if not document_id:
            return {"error": "source='kb_document' requires 'document_id'."}
        resolved = await _fetch_owned_document_bytes(ctx, document_id)
        if isinstance(resolved, dict):  # owner-scope refusal (T-03) → propagate honestly
            return resolved
        _fname, doc_bytes, mime = resolved
        return (doc_bytes, mime)

    return {"error": (
        f"Unknown source '{source}'. Use one of: workspace, sandbox_output, inline, "
        "kb_document."
    )}


async def _handle_attach_skill_file(args: dict, ctx: ToolContext) -> ToolResult:
    """FILE-01 — attach a file onto a skill the caller OWNS, from four sources (D-05).

    G-5: handler + one ``_TOOL_REGISTRY`` line + one ``get_tools()`` schema; threads.py
    untouched, no ``provider ==`` fork. self_improve-gated (D-11) in ``get_tools()`` AND
    refused in-flight via ``_CAPABILITY_FLAG_TOOLS``. Owner-only WRITE gate (D-06/T-04):
    the target skill is resolved by name under ``.eq("user_id")`` — NEVER the
    ``.or_(is_org_shared.eq.true)`` READ filter — and ``is_system`` skills are also rejected;
    service-role has no RLS backstop so this app gate is load-bearing. A colliding filename
    overwrites in place via a race-immune upsert (D-07, Plan 02 unique index). Reuses the
    existing ``skill_files`` table + ``skill-files`` bucket (D-08) — no new table/bucket.
    """
    import os as _os_local
    import re as _re_local

    uid = ctx.current_user["id"]
    target_skill_name = (args.get("target_skill_name") or "").strip()
    filename_raw = (args.get("filename") or "").strip()
    source = (args.get("source") or "").strip()

    if not target_skill_name:
        return ToolResult(result=json.dumps({"error": "target_skill_name is required."}))
    if not filename_raw:
        return ToolResult(result=json.dumps({"error": "filename is required."}))

    # T-02 / Pitfall 4 — sanitize the MODEL-supplied filename to a safe basename so it can
    # never carry a "../" traversal segment out of the owner-prefixed Storage path. For a
    # normal filename this is a no-op (basename + charset scrub, matching FILE-02's stem).
    stem = _os_local.path.basename(filename_raw)
    filename = _re_local.sub(r"[^a-zA-Z0-9._\- ]", "_", stem)
    filename = _re_local.sub(r"\.{2,}", ".", filename).strip() or "attachment"

    # ── D-06 / T-04 / SC#4 — resolve the target skill OWNER-ONLY. Empty .data (not owned
    #    / not found) OR is_system → refuse. NEVER the .or_(is_org_shared.eq.true) READ filter.
    #    WR-04: the gate intentionally does NOT reject is_org_shared — an org-shared skill the
    #    caller OWNS is writable BY DESIGN (T-03: attach-to-owned-skill then owner later toggles it
    #    global is a documented, owner-driven data-movement path, not blocked). The refusal
    #    copy below is therefore scoped to built-in (is_system) skills only, so it never
    #    overstates the enforcement (auditors: owner-scope on user_id is the load-bearing gate). ──
    try:
        skill_res = await aexec(
            ctx.supabase.table("skills").select("id, is_system")
            .eq("name", target_skill_name).eq("user_id", uid).maybe_single()
        )
        skill_row = skill_res.data if skill_res else None
    except Exception:  # noqa: BLE001 — WR-01: honest refusal on a bad-arg/transient DB failure, no raw DB text leak
        skill_row = None
    if not skill_row or skill_row.get("is_system"):
        return ToolResult(result=json.dumps({"error": (
            f"No skill named '{target_skill_name}' that you own — you can only attach "
            "files to a skill you own (built-in skills are never writable)."
        )}))
    skill_id = skill_row["id"]

    # ── Resolve file_bytes + mime by the source discriminator (D-05) ──
    resolved = await _resolve_attach_source_bytes(args, ctx, source, filename)
    if isinstance(resolved, dict):  # honest-failure convention (weak-model guard / T-03)
        return ToolResult(result=json.dumps(resolved))
    file_bytes, mime = resolved

    # ── T-02 — owner-prefixed Storage path. skill_id from the RESOLVED owned skill, uid
    #    from ctx, filename sanitized above — no segment is ever a raw model arg. ──
    storage_path = f"{uid}/{skill_id}/{filename}"

    # ── D-07 pre-check drives the created/updated REPORT only (informational). Write
    #    correctness is the race-immune upsert below (Plan 02 unique index), so even if two
    #    concurrent same-name attaches both read "not present" the DB still collapses to one
    #    row — the report may say "created" twice but the state stays consistent. ──
    exist_res = await aexec(
        ctx.supabase.table("skill_files").select("id")
        .eq("skill_id", skill_id).eq("filename", filename).maybe_single()
    )
    pre_existed = bool(exist_res and exist_res.data)

    # ── D-07 overwrite-in-place — Storage upsert (Pitfall 2: bare .upload() 409s on a
    #    colliding path) + race-immune DB upsert on the (skill_id, filename) unique index.
    #    Both blocking calls are off-loop (Pitfall 1). ──
    uploaded = False
    try:
        await run_in_threadpool(lambda: ctx.supabase.storage.from_("skill-files").upload(
            path=storage_path, file=file_bytes,
            file_options={"content-type": mime, "upsert": "true"},
        ))
        uploaded = True
        await aexec(ctx.supabase.table("skill_files").upsert(
            {
                "skill_id": skill_id, "user_id": uid, "filename": filename,
                "file_path": storage_path, "file_size": len(file_bytes), "mime_type": mime,
            },
            on_conflict="skill_id,filename",
        ))
    except Exception as exc:  # noqa: BLE001 — honest tool-result error, never raise into the loop
        logger.exception(
            "attach_skill_file write failed (skill=%s file=%s)", skill_id, filename
        )
        # IN-01 — the object may already be committed to the bucket while the skill_files
        # upsert failed (e.g. migration 101's unique index not yet applied → ON CONFLICT
        # error 42P10, or any transient DB blip). Best-effort remove the just-uploaded object
        # so a failed attach does not accrete an untracked orphan (the tool hard-depends on
        # migration 101; until it lands on cloud every attach would otherwise leave one behind).
        if uploaded:
            try:
                await run_in_threadpool(
                    lambda: ctx.supabase.storage.from_("skill-files").remove([storage_path])
                )
            except Exception:  # noqa: BLE001 — cleanup is best-effort; never mask the original error
                logger.warning(
                    "attach_skill_file: could not remove orphaned object %s", storage_path
                )
        return ToolResult(result=json.dumps({
            "error": f"Failed to attach the file to the skill: {exc}"
        }))

    status = "updated" if pre_existed else "created"

    # Optional additive SSE — provider-uniform (no provider== fork), best-effort no-op when
    # the run/emit substrate is absent (mirrors workspace_file_written).
    try:
        if getattr(ctx, "emit", None) and getattr(ctx, "run_id", None):
            await ctx.emit(
                ctx.redis, ctx.run_id, "skill_file_attached",
                skill=target_skill_name, filename=filename, status=status,
            )
    except Exception:  # noqa: BLE001 — an emit failure must never break a clean attach
        logger.exception("skill_file_attached emit failed for run %s", getattr(ctx, "run_id", None))

    return ToolResult(result=json.dumps({
        "status": status, "filename": filename, "skill": target_skill_name,
    }))


# Phase 142 (SRH-01 / D-05 / D-05b) — the proactive per-skill runtime note ridden
# along the load_skill RESULT (the load_skill-flag half of the D-05 proactive home).
# Names the non-Python script file(s) a skill bundles that the Python-only sandbox
# cannot run. Model-facing, advisory only — mirrors save_skill's lint_warnings.
_SKILL_RUNTIME_NOTE = (
    "This skill bundles non-Python script file(s) ({names}) that this Python-only "
    "sandbox cannot execute. Their content may still guide you; do not try to run "
    "them as programs."
)


def _skill_runtime_note(file_names: list[str]) -> str | None:
    """Return an advisory note naming any bundled non-Python script files, else None.

    Pure: scans each filename's extension against the shared Plan-01 ``SCRIPT_EXTS``
    allowlist. Non-blocking by contract — the caller computes this defensively and
    NEVER lets it fail a load (the ``lint_warnings`` posture). Returns ``None`` when
    the file list is all-Python / non-script (nothing to warn about).
    """
    offending = [
        name for name in file_names
        if os.path.splitext(name)[1].lstrip(".").lower() in SCRIPT_EXTS
    ]
    if not offending:
        return None
    return _SKILL_RUNTIME_NOTE.format(names=", ".join(offending))


# ─────────────────────────────────────────────────────────────────────────────
# SEED-125 (CR-01) — org-gated skill resolution on the service-role client.
#
# The agent's skill tools (load_skill / read_skill_file / save_skill sibling-lint /
# execute_code skill-file injection) resolve skills on the BYPASSRLS service-role
# producer client (``ctx.supabase``), so the membership org gate that RLS + the mig-110
# DEFINER fns enforce on EVERY request path never applies here. Without an org predicate
# the legacy ``.or_(user_id.eq.<caller>,is_org_shared.eq.true)`` filter matched ANY org's
# ``is_org_shared`` skill → a disjoint-org caller's agent could load another org's skill
# instructions + pull its bundled file bytes. This is the SKILLS analog of the SEED-124
# folder leak Phase 165 closed on the same service-role seam.
#
# Phase 182 (CR-01) HOISTED the predicate itself to ``app.utils.skill_visibility`` — the
# grounding/canvas seam needs the SAME rule but cannot import this module (a real
# harness.scope → task_service → tool_dispatcher cycle), and a second copy is exactly the
# drift the one-source red line forbids. The rule, the RLS shape it mirrors and the
# fail-closed reasoning all live in that module's docstring. This module keeps only the
# ToolContext-shaped resolver below.
# ─────────────────────────────────────────────────────────────────────────────

async def _resolve_skill_visibility_or(ctx: ToolContext, *, born_for: bool = False) -> str:
    """Resolve the caller's org set + build the org-gated skill-visibility ``.or_()``.

    The service-role client bypasses RLS so ``auth.uid()`` / ``current_user_org_ids()``
    never resolve here — the org set MUST be resolved from the threaded ``user_id`` via
    ``org_members`` (fail-closed on an empty membership), exactly as ``folder_utils`` does
    for the folder analog (SEED-124). One await per handler; handlers with two resolution
    sites (read_skill_file / execute_code) resolve ONCE and reuse the returned string so
    the injection loop never fires N membership round-trips.

    ⭐ **Phase 264 (PACK-17 / D-264-04) — ``born_for`` IS THE DECISION, NOT A CONVENIENCE.**
    When true, the caller's ACTIVE Expert bundle (the ``born_for_bundle_id`` field on ``ctx``,
    set only by the agent-loop ToolContext builds off an ACCESS-CHECKED
    ``ResolvedExpertBundle.bundle_id``) is
    forwarded to ``build_skill_visibility_or``, which nests a
    ``born_for_expert_bundle_id = <bundle> AND is_enabled`` disjunct INSIDE the org gate. A
    resolver that simply read ``ctx`` unconditionally would have widened **all four** call
    sites by OMISSION — including ``_handle_save_skill``'s lint corpus, which D-264-04
    deliberately refuses. The explicit keyword makes every site's decision readable at the
    site, and a widening therefore cannot happen by silence. The per-site reasons are written
    beside each of the four calls below; the count is fenced at exactly three opting-in call
    sites by ``tests/unit/test_264_load_skill_born_for.py`` — as an AST count, because a
    comment quoting the keyword satisfies a grep and wires nothing.

    ⛔ The field is read through ``getattr`` with a ``None`` default, never as a bare attribute
    on ``ctx`` (the one read in this module is the line below) — three
    existing suites build duck-typed ``ToolContext``-shaped stubs that predate this field, and
    an ``AttributeError`` here would be swallowed by the ``save_skill`` lint wrapper and the
    ``execute_code`` outer guard rather than surfacing. The same ``getattr`` precedent guards
    ``skill_instructions_override`` and ``skill_snapshot`` in this module.

    ⛔ The bundle is coerced with ``str(...)`` at THIS seam: ``ToolContext`` types the field
    ``UUID | None`` while ``app.utils.skill_visibility`` annotates its parameter ``str | None``
    (264-02's zero-new-imports convention — that module imports nothing but ``coerce_uid``).
    Never re-derive the predicate term here; pass the bundle THROUGH, or the one-home count
    fence in ``tests/unit/test_264_one_home_born_for_predicate.py`` names this module.
    """
    org_ids = await _resolve_caller_org_ids(ctx.supabase, ctx.current_user["id"])
    _bundle = getattr(ctx, "born_for_bundle_id", None) if born_for else None
    return build_skill_visibility_or(
        ctx.current_user["id"],
        org_ids,
        expert_bundle_id=str(_bundle) if _bundle is not None else None,
    )


async def _handle_load_skill(args: dict, ctx: ToolContext) -> ToolResult:
    skill_name = args.get("skill_name", "")
    # Emit skill_activated SSE event immediately (SKIL-12)
    await ctx.emit(ctx.redis, ctx.run_id, 'skill_activated', skill_name=skill_name)
    # Resolve skill -- on a name collision the most-authoritative row wins:
    # system > global > owned (SEED-102). is_system DESC pins a protected built-in
    # above any same-named owned row; is_org_shared DESC is the secondary tie-break.
    # SEED-125 (CR-01): the visibility filter is org-gated (is_system universal escape
    # OR org_id ∈ caller_org_ids AND (owner OR is_org_shared)) — a disjoint-org caller no
    # longer resolves another org's is_org_shared skill on the BYPASSRLS service client.
    #
    # Phase 264 (PACK-17 / D-264-04) — ⭐ WIDEN. THIS SITE IS THE DEFECT. `load_skill` is
    # advertised to every chat run, Expert or not (Phase 267 deleted the Expert tool filter,
    # D-267-01), and Phase 263 already admits a colleague's born-for skill at
    # RESOLVE time — but that result is the CATALOG only (names + descriptions). The
    # instruction BODY is fetched here, through a predicate that had never heard of
    # `born_for_expert_bundle_id`, so for every org member but the author the prompt promised
    # a skill the load path then refused. The opt-in below closes exactly that.
    # ⛔ Resolved ONCE and reused at the miss branch below, which is what keeps the
    # `available_skills` listing built from the SAME filter as the primary query — it can
    # never name a set the query would not admit (RESEARCH §8.8 / T-264-14). Re-deriving the
    # filter there would silently re-open half the defect while every outcome test stayed green.
    _skill_filter = await _resolve_skill_visibility_or(ctx, born_for=True)
    _skill_resp = await aexec(
        ctx.supabase.table("skills")
        .select("id, name, description, instructions, user_id")
        .or_(_skill_filter)
        .eq("name", skill_name)
        .eq("is_enabled", True)
        .order("is_system", desc=True).order("is_org_shared", desc=True)
    )
    skill_row = _skill_resp.data

    if not skill_row:
        # Exact-name miss. Weaker models -- local ones especially -- emit the skill's
        # HUMAN-READABLE title ("Weekly Report Writer") where the registry stores a slug
        # ("weekly-report-writer"). Measured 2026-08-18 on openai/gpt-oss-20b: the run
        # completed but told the operator to upload a template that WAS already attached,
        # because the miss above returned a dead end -- an error naming no valid
        # alternative, so the model could not self-correct and reasoned on from a false
        # premise. A stronger cloud model emits the slug first try and never reaches this
        # branch, which is exactly why the gap read as "local models are broken".
        #
        # Two additive recoveries, both reached ONLY where the code above already failed:
        #   1. normalised match (casefold, separators unified) -- resolves the title form
        #   2. an error that LISTS the loadable names, so one retry can succeed
        import re as _re_local   # module-local idiom used elsewhere in this file

        def _norm(v: str) -> str:
            return _re_local.sub(r"[\s_-]+", "-", (v or "").strip().casefold())

        _all_resp = await aexec(
            ctx.supabase.table("skills")
            .select("id, name, description, instructions, user_id")
            .or_(_skill_filter)
            .eq("is_enabled", True)
            .order("is_system", desc=True).order("is_org_shared", desc=True)
        )
        _candidates = _all_resp.data or []
        if not isinstance(_candidates, list):
            _candidates = [_candidates]

        _target = _norm(skill_name)
        # First match wins: the query keeps the is_system > is_org_shared precedence
        # the exact-match path relies on (SEED-102 / SEED-125), so iteration order IS
        # the authority order. Never re-sort here.
        _match = next((c for c in _candidates if _norm(c.get("name")) == _target), None)

        if _match is None:
            _available = sorted({c.get("name") for c in _candidates if c.get("name")})
            return ToolResult(result=json.dumps({
                "error": f"Skill '{skill_name}' not found or not enabled.",
                "available_skills": _available,
                "hint": "Call load_skill again with one of the names in available_skills, exactly as written.",
            }))

        logger.info(
            "load_skill: resolved '%s' to '%s' by normalised name", skill_name, _match.get("name"),
        )
        skill_row = _match

    row = skill_row[0] if isinstance(skill_row, list) else skill_row

    # Phase 067.1 Plan 04: follow-up emit with skill description
    if row.get("description"):
        await ctx.emit(
            ctx.redis,
            ctx.run_id,
            'skill_loaded',
            skill_name=skill_name,
            description=row["description"],
        )
    ctx.spawn(write_audit_entry(
        user_id=ctx.current_user["id"],
        action_type="skill.load",
        metadata={"skill_id": row["id"], "skill_name": row["name"]},
        supabase=ctx.supabase,
    ))
    # Fetch attached filenames (FILE-04)
    _files_resp = await aexec(
        ctx.supabase.table("skill_files")
        .select("filename")
        .eq("skill_id", row["id"])
        .order("filename")
    )
    files_data = _files_resp.data or []
    file_names = [f["filename"] for f in files_data]
    # Phase 135 (135-02 / SI-01) — Pitfall #1 fix: return the DRAFT instructions
    # when the re-eval passed a skill_instructions_override map containing THIS
    # skill's name; else the live DB row's body (byte-identical Deep). getattr so a
    # duck-typed ctx stub predating the field still works (096 workflow_run_id
    # precedent). Keyed on skill_name — the SAME value `.eq("name", ...)` looked up.
    # 2026-08-18: the normalised-name fallback above means `skill_name` (what the MODEL
    # typed) may differ from `row["name"]` (what actually resolved). This map is keyed by
    # the REAL skill name, so check the resolved name too -- checking only the caller's
    # string would silently serve LIVE instructions to a re-eval that asked for DRAFT
    # ones, and the eval would score the wrong text while reporting success.
    override = getattr(ctx, "skill_instructions_override", None)
    instructions = row["instructions"]
    if override is not None:
        for _key in (skill_name, row.get("name")):
            if _key is not None and _key in override:
                instructions = override[_key]
                break
    result_payload = {
        "name": row["name"],
        "instructions": instructions,
        "files": file_names,
    }
    # D-05b — attach a non-blocking runtime note when the skill bundles a script the
    # Python-only sandbox cannot execute. Computed defensively (the save_skill
    # lint_warnings posture): any failure degrades to no note, and the key is added
    # ONLY when present so the all-Python result stays byte-identical. Never blocks.
    try:
        runtime_note = _skill_runtime_note(file_names)
    except Exception:
        runtime_note = None
    if runtime_note is not None:
        result_payload["runtime_note"] = runtime_note
    return ToolResult(result=json.dumps(result_payload))


async def _handle_save_skill(args: dict, ctx: ToolContext) -> ToolResult:
    name = args.get("name", "").strip()
    description = args.get("description", "")
    instructions = args.get("instructions", "")
    if not name:
        return ToolResult(result=json.dumps({"error": "Skill name is required."}))

    # Check if user already owns a skill with this name
    existing_resp = await aexec(
        ctx.supabase.table("skills")
        .select("id")
        .eq("user_id", ctx.current_user["id"])
        .eq("name", name)
        .limit(1)
    )
    existing = existing_resp.data[0] if existing_resp.data else None

    # TRIG-03 (D-09/D-10/Pitfall 6): lint the description, owner-scoped, NEVER block.
    # The save always proceeds; warnings ride along as a non-fatal note so the agent
    # can mention them. A read failure degrades to an empty sibling list.
    lint_warnings: list[dict] = []
    try:
        # SEED-125 (CR-01): org-gate the sibling set so the lint never reads (or echoes
        # the description of) another org's is_org_shared skill on the service client.
        #
        # Phase 264 (PACK-17 / D-264-04) — ⛔ DELIBERATELY **NOT** WIDENED. This is the one
        # resolver call in this module that keeps the default `born_for=False`, and the
        # decision is written here rather than left as an omission. Three measured reasons:
        #   1. ⚠ CORRECTED BY PHASE 267 — the original reason #1 was "`save_skill` is in NEITHER
        #      `EXPERT_CORE_TOOLS` nor `EXPERT_DELIVERABLE_TOOLS`, so it is **not advertised** to
        #      an Expert run". That became FALSE when 267 deleted both constants and the Expert
        #      tool filter (D-267-01): `save_skill` IS advertised to Expert threads now. The
        #      decision NOT to widen still stands on reasons 2 and 3 alone, which never depended
        #      on it. (Chat still has no dispatch-side whitelist — `ctx.phase_whitelist` is `None`
        #      on a chat run; that backstop is owed to SEED-303, not built.)
        #   2. The read feeds a non-blocking description **lint** corpus and produces NO
        #      user-visible capability. Widening it would only let an Expert's borrowed skills
        #      influence another user's save-time warnings — a widening with no upside.
        #   3. It is a WRITE handler's helper, while PACK-17's axis is *read the body you were
        #      promised*.
        # Driven, not merely asserted: `tests/unit/test_264_unchanged_sites_fenced.py` resolves
        # this filter on a ctx that DOES carry a bundle and pins it to the base literal.
        _sibling_filter = await _resolve_skill_visibility_or(ctx)
        siblings_resp = await aexec(
            ctx.supabase.table("skills")
            .select("id, description")
            .or_(_sibling_filter)
        )
        siblings = [
            r.get("description", "")
            for r in (siblings_resp.data or [])
            if not existing or str(r.get("id")) != str(existing["id"])
        ]
        lint_warnings = lint_description(name, description, siblings)
    except Exception:
        lint_warnings = []

    if existing:
        row = existing
        await aexec(
            ctx.supabase.table("skills").update({
                "description": description,
                "instructions": instructions,
            }).eq("id", row["id"]).eq("user_id", ctx.current_user["id"])
        )
        return ToolResult(result=json.dumps(
            {"status": "updated", "name": name, "lint_warnings": lint_warnings}
        ))
    else:
        await aexec(
            ctx.supabase.table("skills").insert({
                "user_id": ctx.current_user["id"],
                "name": name,
                "description": description,
                "instructions": instructions,
            })
        )
        return ToolResult(result=json.dumps(
            {"status": "created", "name": name, "lint_warnings": lint_warnings}
        ))


# Phase 142 (SRH-01 / SC#3 / D-11) — the honest caveat prepended to a decoded
# non-Python script (an ext in the shared SCRIPT_EXTS). read_skill_file stops
# mislabeling script text as unreadable "binary"; the model is told the file is
# reference-only, not runnable here. PLAIN TEXT (not json) — matches the .py/.md
# text-return contract so the model reads it as source, not an error object.
_SCRIPT_REF_CAVEAT = (
    "[reference only — '{filename}' is a {ext} script; this sandbox runs Python only "
    "and cannot execute it. Read it for reference; do not attempt to run it.]\n\n"
)


def _decode_skill_file_bytes(filename: str, raw_bytes: bytes) -> str:
    """Decode skill-file bytes to a text tool_result by extension.

    PURE EXTRACTION (099 Plan 03) of the ext-decode block that lived inline in
    ``_handle_read_skill_file`` — byte-identical behavior (docx/xlsx/pptx/text/binary).
    Shared by BOTH the live-skill read path (the SC#3 red line — unchanged behavior)
    AND the 099 snapshot-routing branch, so the snapshot read decodes exactly as the
    live read does. Behavior here MUST stay identical to the pre-099 inline block.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "docx":
        import docx as _docx  # python-docx
        doc = _docx.Document(io.BytesIO(raw_bytes))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    elif ext == "xlsx":
        import openpyxl as _openpyxl
        wb = _openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            for row_data in sheet.iter_rows(values_only=True):
                line = "\t".join(str(c) if c is not None else "" for c in row_data)
                if line.strip():
                    rows.append(line)
        return "\n".join(rows)
    elif ext == "pptx":
        from pptx import Presentation as _Presentation  # python-pptx
        prs = _Presentation(io.BytesIO(raw_bytes))
        slides = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slides.append(shape.text)
        return "\n".join(slides)
    elif ext in {"txt", "md", "py", "csv", "json", "yaml", "yml", "toml", "html", "xml", "rst", "log"}:
        return raw_bytes.decode("utf-8", errors="replace").replace('\x00', '')
    elif ext in SCRIPT_EXTS:
        # SC#3 / D-11 — non-Python scripts (.js/.sh/...) ARE text: decode them as
        # honest reference source with a "not executable in this sandbox" caveat,
        # replacing the misleading "binary — upload a text version" else-branch below.
        # This lives in the SHARED decoder, so the live read (_handle_read_skill_file
        # live path) and the 099 snapshot read inherit the identical string for free
        # (byte-symmetry / Pitfall 3). SCRIPT_EXTS is the Plan-01 single source.
        source = raw_bytes.decode("utf-8", errors="replace").replace('\x00', '')
        return _SCRIPT_REF_CAVEAT.format(filename=filename, ext=ext) + source
    else:
        # Unrecognized or binary type
        return json.dumps({
            "error": f"File '{filename}' is a binary file that cannot be read as text. "
                     "Upload a text-based version instead."
        })


async def _handle_read_skill_file(args: dict, ctx: ToolContext) -> ToolResult:
    filename = args.get("filename", "")
    # 099 D-04 GATE (mirrors the 098 search-scope gate at :187) — when a workflow
    # phase carries a materialized skill snapshot, read from the IMMUTABLE snapshot
    # copies, NOT the live skill. None (Deep mode + non-skill phases) => this branch
    # is skipped and the live path below runs BYTE-IDENTICAL (SC#3 red line / Pitfall 4).
    snapshot = getattr(ctx, "skill_snapshot", None)
    if snapshot is not None:
        if filename not in getattr(snapshot, "files", []):
            return ToolResult(result=json.dumps(
                {"error": f"File '{filename}' not in the workflow's skill snapshot."}
            ))
        storage_path = f"{snapshot.storage_prefix}/{filename}"
        try:
            # Un-wrapped .download() for byte-symmetry with the live path (Open Question 4 —
            # single small file; only the multi-file WRITE materializer is threadpool-wrapped).
            raw_bytes = ctx.supabase.storage.from_("skill-files").download(storage_path)
            tool_result = _decode_skill_file_bytes(filename, raw_bytes)
        except Exception as e:
            tool_result = json.dumps({"error": f"File '{filename}' not found in snapshot: {e}"})
        return ToolResult(result=tool_result)

    # ── live-skill resolution below — SC#3 red line preserved (Pitfall 4); only the
    #    visibility filter is org-gated per SEED-125 (CR-01). Resolve the caller's org
    #    set ONCE and reuse it for the normalized-name retry (no double round-trip). ──
    skill_name = args.get("skill_name", "")
    # Phase 264 (PACK-17 / D-264-04) — ⭐ WIDEN. `read_skill_file` is advertised to every chat run,
    # and `load_skill` returns `files: [...]` (:1405) — so a body that loads while its bundled
    # files 404 is the SAME defect one layer down: the model is told the files exist and then
    # cannot read them. ⚠ This handler applies NO `is_enabled` filter of its own, which is why
    # the born-for disjunct carries its own `is_enabled.is.true` term (264-02, Form 2 /
    # T-264-15) — a bare arm here would admit a DISABLED born-for skill's bundled bytes.
    # Resolved ONCE and reused by the normalised-name retry below.
    _skill_filter = await _resolve_skill_visibility_or(ctx, born_for=True)
    # Resolve skill to get owner's user_id for storage path
    _sr_resp = await aexec(
        ctx.supabase.table("skills")
        .select("id, user_id")
        .or_(_skill_filter)
        .eq("name", skill_name)
        .maybe_single()
    )
    skill_row = _sr_resp.data if _sr_resp is not None else None
    if not skill_row:
        # Retry with normalized name for agent display-name mismatches
        _sr_norm = skill_name.lower().replace(" ", "-")
        if _sr_norm != skill_name:
            _sr_resp2 = await aexec(
                ctx.supabase.table("skills")
                .select("id, user_id")
                .or_(_skill_filter)
                .eq("name", _sr_norm)
                .maybe_single()
            )
            skill_row = _sr_resp2.data if _sr_resp2 is not None else None
    if not skill_row:
        return ToolResult(result=json.dumps({"error": f"Skill '{skill_name}' not found."}))

    row = skill_row[0] if isinstance(skill_row, list) else skill_row
    storage_path = f"{row['user_id']}/{row['id']}/{filename}"
    try:
        raw_bytes = ctx.supabase.storage.from_("skill-files").download(storage_path)
        tool_result = _decode_skill_file_bytes(filename, raw_bytes)
    except Exception as e:
        tool_result = json.dumps({"error": f"File '{filename}' not found: {e}"})

    return ToolResult(result=tool_result)
