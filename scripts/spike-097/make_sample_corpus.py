"""
THROWAWAY (Phase 097 spike, SEED-051) — generate a small, realistic risk-content
corpus for the dev test user to ingest into a fresh KB folder.

Why: the spike needs genuine risk material (causes / events / effects / owners /
mitigations / status) to ground the template-fill experiment (097-02) on. None of
the existing test folders matched the risk heuristic, so we author a controlled
fictional project ("Project Meridian") spread across THREE document types a real PM
would have — a weekly status report (prose), risk workshop notes (semi-structured),
and a project charter excerpt (strategic). Risks deliberately recur across docs so
citation cross-referencing is exercised.

Output: scripts/spike-097/sample-corpus/*.docx  (upload these via the UI).
Run from repo root:  backend/venv/Scripts/python.exe scripts/spike-097/make_sample_corpus.py
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).parent / "sample-corpus"
OUT.mkdir(exist_ok=True)


def _title(doc, text):
    h = doc.add_heading(text, level=0)
    return h


def _meta(doc, pairs):
    for k, v in pairs:
        p = doc.add_paragraph()
        r = p.add_run(f"{k}: ")
        r.bold = True
        p.add_run(v)


# ---------------------------------------------------------------------------
# Doc 1 — Weekly Status Report (prose; risks woven into a narrative section)
# ---------------------------------------------------------------------------
def status_report():
    doc = Document()
    _title(doc, "Project Meridian — Weekly Status Report")
    _meta(doc, [
        ("Project", "Meridian (Legacy CRM → Cloud Platform Migration)"),
        ("Reporting period", "Week 9 of 24"),
        ("Prepared by", "Dana Whitfield, Project Manager"),
        ("Overall RAG status", "AMBER"),
    ])

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "Project Meridian is migrating Brightline Insurance's on-premise legacy CRM "
        "to a cloud-hosted platform over a 24-week programme. We are at the end of "
        "Week 9. Foundation infrastructure is provisioned and the data-profiling pass "
        "is complete, but two risks have escalated this week and the programme RAG "
        "status has moved from Green to Amber. The critical path now runs through the "
        "data-migration validation gate and the third-party policy-rating API "
        "integration."
    )

    doc.add_heading("2. Progress This Week", level=1)
    for line in [
        "Cloud landing zone provisioned and security-hardened (sign-off received).",
        "Legacy data profiling complete: 412,000 customer records, 38,000 policies.",
        "Identified ~6% duplicate customer records in the legacy CRM during profiling.",
        "Drafted the cutover runbook (v0.3) for review at the next steering committee.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("3. Upcoming Milestones", level=1)
    for line in [
        "Week 11 — Data migration dress rehearsal (dependent on dedup completion).",
        "Week 13 — Integration test entry for the policy-rating API.",
        "Week 18 — User Acceptance Testing begins.",
        "Week 24 — Production cutover and legacy decommission.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("4. Key Risks and Issues", level=1)
    doc.add_paragraph(
        "Risk M-01 (Data quality). Profiling found roughly 6% duplicate and partially "
        "merged customer records in the legacy CRM. If these are migrated as-is, the "
        "event is corrupted or duplicated customer records in the new platform, with "
        "the effect of mis-routed policy correspondence and a loss of customer trust. "
        "Likelihood is assessed High and impact High. Owner: Priya Nair (Data Lead). "
        "Mitigation: a deterministic de-duplication pass plus a validation gate that "
        "blocks migration if the duplicate rate exceeds 0.5%. Status: in progress."
    )
    doc.add_paragraph(
        "Risk M-02 (Integration dependency). The contract for the third-party "
        "policy-rating API has not been finalised by the vendor, Acme Ratings Ltd. The "
        "event is that integration testing in Week 13 cannot start on time; the effect "
        "is a direct slip to the UAT entry date. Likelihood Medium, impact High. "
        "Owner: Tom Becker (Integration Lead). Mitigation: escalate to vendor account "
        "management and prepare a stub-based fallback so internal testing is not "
        "blocked. Status: open and escalating."
    )
    doc.add_paragraph(
        "Risk M-03 (Resourcing). Two senior engineers are currently shared 50/50 with "
        "the parallel Billing programme. The cause is an unresolved allocation "
        "conflict; the event is reduced migration throughput; the effect is a "
        "potential two-week schedule slip. Likelihood Medium, impact Medium. Owner: "
        "Dana Whitfield (PM). Mitigation: a dedicated-allocation request has been "
        "submitted to the resourcing board. Status: mitigating."
    )

    doc.add_heading("5. Decisions Needed", level=1)
    doc.add_paragraph(
        "Steering committee approval is requested for the dedicated allocation of the "
        "two senior engineers (see Risk M-03) and for the stub-based integration "
        "fallback budget (see Risk M-02)."
    )
    path = OUT / "Project-Meridian-Status-Report-Week09.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# Doc 2 — Risk Workshop Notes (semi-structured risk log)
# ---------------------------------------------------------------------------
def workshop_notes():
    doc = Document()
    _title(doc, "Project Meridian — Risk Workshop Notes")
    _meta(doc, [
        ("Session", "Fortnightly Risk Review #4"),
        ("Date", "Week 9, Thursday"),
        ("Facilitator", "Dana Whitfield (PM)"),
        ("Attendees", "P. Nair (Data), T. Becker (Integration), S. Okafor (Security), "
                      "R. Lindqvist (Change), M. Alvarez (Sponsor delegate)"),
    ])

    doc.add_paragraph(
        "The team reviewed the open risk register and identified new risks. Each entry "
        "below records cause, event, effect, likelihood, impact, owner and the proposed "
        "response strategy (avoid / reduce / transfer / accept)."
    )

    risks = [
        ("M-01", "Data quality",
         "Legacy CRM contains ~6% duplicate/merged customer records (carried from Week 8).",
         "Duplicate or corrupted records are migrated into the cloud platform.",
         "Mis-routed correspondence, compliance exposure, rework.",
         "High", "High", "Priya Nair (Data Lead)",
         "Reduce — dedup pass + a hard validation gate at 0.5% duplicate threshold.",
         "In progress"),
        ("M-02", "Integration dependency",
         "Policy-rating API contract with Acme Ratings Ltd not finalised.",
         "Integration test entry (Week 13) cannot begin on schedule.",
         "UAT start slips; downstream cutover date at risk.",
         "Medium", "High", "Tom Becker (Integration Lead)",
         "Reduce — vendor escalation + stub-based fallback for internal testing.",
         "Open"),
        ("M-04", "Security / compliance (GDPR)",
         "Production PII is copied into lower test environments for realistic UAT.",
         "A data breach or unauthorised access in a less-controlled environment.",
         "Regulatory fine, reputational damage, mandatory breach notification.",
         "Medium", "High", "Sade Okafor (Security Lead)",
         "Reduce — anonymise/pseudonymise all PII before it leaves production.",
         "Open"),
        ("M-05", "Cutover / rollback",
         "The Week-24 production cutover has a narrow weekend window and no rehearsed rollback.",
         "Cutover fails partway and the team cannot cleanly revert to legacy.",
         "Extended outage of customer-facing CRM during business hours.",
         "Low", "High", "Tom Becker (Integration Lead)",
         "Reduce — add a Week-22 rollback rehearsal and a go/no-go checkpoint.",
         "Open"),
        ("M-06", "User adoption / training",
         "Front-line agents have had no exposure to the new platform UI.",
         "Agents struggle at go-live and revert to manual workarounds.",
         "Productivity dip, data captured outside the system, support load spike.",
         "Medium", "Medium", "Rita Lindqvist (Change Lead)",
         "Reduce — phased training from Week 16 plus floor-walker support at go-live.",
         "Not started"),
        ("M-07", "Budget / scope",
         "Stakeholders are requesting additional custom reports beyond the agreed scope.",
         "Uncontrolled scope creep consumes contingency budget.",
         "Budget overrun and distraction from critical-path migration work.",
         "Medium", "Medium", "Dana Whitfield (PM)",
         "Reduce — route all new requests through formal change control.",
         "Mitigating"),
    ]

    cols = ["Risk ID", "Category", "Cause", "Event", "Effect",
            "Likelihood", "Impact", "Owner", "Response strategy", "Status"]
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Light Grid Accent 1"
    for i, c in enumerate(cols):
        run = table.rows[0].cells[i].paragraphs[0].add_run(c)
        run.bold = True
    for row in risks:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val

    doc.add_paragraph()
    doc.add_paragraph(
        "Action: M-05 rollback rehearsal to be added to the plan; M-04 anonymisation "
        "approach to be confirmed with the Data Protection Officer before Week 12."
    )
    path = OUT / "Project-Meridian-Risk-Workshop-Notes.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# Doc 3 — Project Charter excerpt (strategic risks)
# ---------------------------------------------------------------------------
def charter():
    doc = Document()
    _title(doc, "Project Meridian — Project Charter (Excerpt)")
    _meta(doc, [
        ("Document", "Project Charter v1.2 — Section 7 extract"),
        ("Sponsor", "Marcus Alvarez, Chief Operating Officer"),
        ("Programme", "CRM Modernisation"),
    ])

    doc.add_heading("7. Initial Risk Assessment", level=1)
    doc.add_paragraph(
        "At charter approval the steering committee recorded the following strategic "
        "risks to be actively managed by the programme. These are reviewed at each "
        "stage gate."
    )
    doc.add_paragraph(
        "Strategic risk SR-01 (Executive sponsorship continuity). The cause is a "
        "pending reorganisation of the operations division. The event is loss of an "
        "engaged executive sponsor mid-programme; the effect is stalled decisions and "
        "funding uncertainty. Owner: Marcus Alvarez (COO). Response: name a deputy "
        "sponsor with delegated authority. Status: open."
    )
    doc.add_paragraph(
        "Strategic risk SR-02 (Regulatory approval). Brightline operates in a regulated "
        "insurance market; the cloud platform must pass an information-security "
        "assessment before holding customer data. If approval is delayed, the event is "
        "a blocked production cutover, with the effect of a schedule slip beyond the "
        "24-week window. Likelihood Low, impact High. Owner: Sade Okafor (Security "
        "Lead). Response: engage the regulator early and book the assessment in "
        "Week 14. Status: open."
    )
    doc.add_paragraph(
        "Strategic risk SR-03 (Legacy decommission dependency). The legacy CRM also "
        "feeds the nightly actuarial reporting batch. The cause is a hidden downstream "
        "dependency; the event is that decommissioning legacy breaks actuarial "
        "reporting; the effect is a gap in regulatory financial reporting. Likelihood "
        "Medium, impact High. Owner: Priya Nair (Data Lead). Response: map and "
        "re-point the downstream feed before decommission. Status: open."
    )
    path = OUT / "Project-Meridian-Charter-Excerpt.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    made = [status_report(), workshop_notes(), charter()]
    print("Generated risk corpus:")
    for p in made:
        # quick integrity re-open + rough word count
        d = Document(p)
        words = sum(len(par.text.split()) for par in d.paragraphs)
        for t in d.tables:
            for r in t.rows:
                for c in r.cells:
                    words += len(c.text.split())
        print(f"  - {p.name}  ({words} words, re-opened OK)")
    print(f"\nUpload these 3 files via the UI into a NEW folder for fhdmrd@gmail.com.")
