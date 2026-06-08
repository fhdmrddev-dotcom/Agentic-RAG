"""Generate the {%tr %} variable-row risk-register.docx template (THROWAWAY spike).

Phase 097 Wave 0 (SEED-051). Builds templates/risk-register.docx programmatically
with python-docx (no manual Word authoring) so the spike has a deterministic,
reproducible template under test. The template carries docxtpl Jinja tags:

  - Two scalar tags in title paragraphs:  {{ project_name.value }} / {{ report_date.value }}
  - A header row + a {%tr %}-driven repeating body row. docxtpl's {%tr %} tag
    REPLACES the entire <w:tr> row that holds it with the bare {% ... %} jinja
    statement (verified in docxtpl/template.py patch_xml, line ~185). So the loop
    is authored as THREE rows:
        row 1: a dedicated {%tr for r in rows %} row  -> becomes {% for r in rows %}
        row 2: the content row {{ r.<field> }}        -> the row that repeats
        row 3: a dedicated {%tr endfor %} row         -> becomes {% endfor %}
    so the content row grows once per `r` in `rows` (the variable-length-rows
    answer to RESEARCH Pitfall 5 / A5).

Each {%tr %} tag is written via `cell.text = ...`, which places it inside a single
run/paragraph — never split across runs (avoids docxtpl Pitfall 4 TemplateSyntaxError).

Run from the repo root:
    backend/venv/Scripts/python.exe scripts/spike-097/make_template.py
"""

from pathlib import Path

from docx import Document

OUT_PATH = Path(__file__).resolve().parent / "templates" / "risk-register.docx"

# Header row (human labels) — column order is the contract.
HEADERS = [
    "Risk ID", "Cause", "Event", "Effect",
    "P", "I", "Score", "Response", "Owner", "Status",
]

# Body row Jinja tags in the SAME column order. Each pulls r.<field>.value from a
# cited field-map row, EXCEPT score, which is a render-time COMPUTE (P x I) and so
# carries no .value (RESEARCH §"The Cited Field-Map Shape").
BODY_TAGS = [
    "{{ r.risk_id.value }}",
    "{{ r.cause.value }}",
    "{{ r.event.value }}",
    "{{ r.effect.value }}",
    "{{ r.probability.value }}",
    "{{ r.impact.value }}",
    "{{ r.score }}",                 # render-time compute — NO .value
    "{{ r.response_strategy.value }}",
    "{{ r.owner.value }}",
    "{{ r.status.value }}",
]

assert len(HEADERS) == len(BODY_TAGS), "header/body column count mismatch"


def build() -> Path:
    doc = Document()

    doc.add_heading("Risk Register", level=1)
    # Scalar tags — each in its own single-run paragraph.
    doc.add_paragraph("Project: {{ project_name.value }}")
    doc.add_paragraph("Date: {{ report_date.value }}")
    doc.add_paragraph("")  # spacer

    # 4 rows: header + a dedicated {%tr for%} row + the content row + a dedicated
    # {%tr endfor%} row. docxtpl replaces each {%tr %}-bearing row with the bare
    # {% ... %} statement, so the loop tags MUST sit in their own rows (not inline
    # in the content row) or jinja sees an unbalanced for/endfor.
    table = doc.add_table(rows=4, cols=len(HEADERS))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for i, label in enumerate(HEADERS):
        header_cells[i].text = label

    # Dedicated loop-open row — tag in the first cell, rest left blank.
    table.rows[1].cells[0].text = "{%tr for r in rows %}"

    # Content row — the row docxtpl repeats once per `r`.
    content_cells = table.rows[2].cells
    for i, tag in enumerate(BODY_TAGS):
        content_cells[i].text = tag

    # Dedicated loop-close row.
    table.rows[3].cells[0].text = "{%tr endfor %}"

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_PATH))
    return OUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"wrote {path}")
