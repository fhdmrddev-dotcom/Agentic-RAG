"""docxtpl render (SSTI-contained, autoescaped) + integrity re-open + corruption log.

Phase 097 Wave 2 (SEED-051) — THROWAWAY spike, Plan 097-03 Task 1. The render half
of the headline end-to-end artifact (SC#1): it turns a cited field-map into an
openable `.docx`, re-opens it as the mandatory corruption guard, scans for residual
placeholders, and writes the named-failure-mode log.

Golden rule (Pattern 1, RESEARCH "## Pattern 1"): the LLM produces DATA only (the
field-map from Plan 02). This module — pinned library code — produces the FILE. The
LLM never touches OOXML bytes; only the field-map context dict flows into docxtpl.

Render is LOCAL (backend venv) here for the fast spike loop. PRODUCTION render is the
sealed, network-less Docker `llm_sandbox` image (Phase 101 — RESEARCH Open Question 3
+ Architectural Responsibility Map). `SandboxedEnvironment` is wired even on this
trusted spike template to PROVE the TMPL-03 SSTI-containment mechanism Phase 101
inherits (threat T-097-08).

Red line / G-5: throwaway code under scripts/spike-097/ only. No backend/app/** edits.
This module is self-contained (docxtpl + python-docx + jinja2 only) so Task 1 can
import it without the backend/dotenv bootstrap.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate
from jinja2 import TemplateSyntaxError
from jinja2.sandbox import SandboxedEnvironment

SPIKE_DIR = Path(__file__).resolve().parent
OUT_DIR = SPIKE_DIR / "out"
CORRUPTION_LOG = OUT_DIR / "corruption.log"

# The nine register columns + the two scalars — used by the residual-placeholder scan
# and documented so the context builder stays in lock-step with templates/risk-register.docx.
ROW_FIELDS = [
    "risk_id", "cause", "event", "effect", "probability",
    "impact", "response_strategy", "owner", "status",
]


# ---------------------------------------------------------------------------
# Context builder — field-map dict -> docxtpl context
# ---------------------------------------------------------------------------
def _num(cited: dict | None):
    """int(value) when the Cited value parses as an int, else None.

    The KB often expresses probability/impact as WORDS (High/Medium/Low) — those do
    NOT parse as ints, so score is left blank. That is an honest, expected result, not
    a bug: the deterministic P x I compute only fires when both are numeric.
    """
    if not cited:
        return None
    v = cited.get("value")
    if v is None:
        return None
    try:
        return int(str(v).strip())
    except (ValueError, TypeError):
        return None


def _cell(cited: dict | None) -> dict:
    """A display-safe Cited dict: value None -> '' so a not-found-in-KB field renders
    as a BLANK cell, never the literal string 'None'. The nested dict shape is kept so
    the template's `{{ r.<field>.value }}` dict-item lookup still resolves."""
    cited = cited or {}
    v = cited.get("value")
    return {
        "value": "" if v is None else v,
        "source_chunk_id": cited.get("source_chunk_id"),
        "source_doc": cited.get("source_doc"),
        "source_page": cited.get("source_page"),
    }


def build_context(field_map_dict: dict) -> dict:
    """Turn a RiskRegisterFieldMap dict (model_dump) into the docxtpl context.

    Returns {"project_name": <cell>, "report_date": <cell>, "rows": [<row>, ...]}.
    Each row carries the nine Cited columns PLUS a render-time `score`:
    `score = int(probability.value) * int(impact.value)` when BOTH parse as ints, else
    blank (the deterministic compute that is NOT an LLM field — RESEARCH §"The Cited
    Field-Map Shape"). None leaves are blanked to '' for clean cells.
    """
    ctx: dict = {
        "project_name": _cell(field_map_dict.get("project_name")),
        "report_date": _cell(field_map_dict.get("report_date")),
        "rows": [],
    }
    for row in field_map_dict.get("rows", []):
        r = {field: _cell(row.get(field)) for field in ROW_FIELDS}
        p, i = _num(row.get("probability")), _num(row.get("impact"))
        score = p * i if (p is not None and i is not None) else None
        r["score"] = "" if score is None else score   # blank cell when not numeric
        ctx["rows"].append(r)
    return ctx


# ---------------------------------------------------------------------------
# Render — EXACTLY the RESEARCH "## Code Examples" shape
# ---------------------------------------------------------------------------
def render(template_path: str, context: dict, out_path: str) -> dict:
    """docxtpl render with SSTI containment + autoescape, then save.

    `SandboxedEnvironment(autoescape=True)` is load-bearing:
      - SandboxedEnvironment  -> SSTI containment (TMPL-03 / threat T-097-08): blocks
        attribute/builtin access if a template were ever malicious. Proven here on the
        trusted template so Phase 101's untrusted-upload path inherits the mechanism.
      - autoescape=True        -> XML-safe values: `& < >` in KB-derived strings are
        escaped to `&amp; &lt; &gt;` so they cannot corrupt the OOXML (Pitfall 2).

    docxtpl owns the bytes; the LLM never does. A `{%tr %}`/`{%p %}` tag split across a
    structural boundary raises TemplateSyntaxError (Pitfall 4) — caught and reported,
    not crashed.

    Returns {"rendered": bool, "error": str | None}.
    """
    doc = DocxTemplate(template_path)
    jenv = SandboxedEnvironment(autoescape=True)   # SSTI containment + XML-safe (&<>)
    try:
        doc.render(context, jinja_env=jenv)        # docxtpl owns the bytes; the LLM never does
    except TemplateSyntaxError as exc:             # Pitfall 4 — tag spans a structural boundary
        return {"rendered": False, "error": f"TemplateSyntaxError: {exc}"}
    doc.save(out_path)
    return {"rendered": True, "error": None}


# ---------------------------------------------------------------------------
# Integrity re-open (the mandatory corruption guard) + residual-tag scan
# ---------------------------------------------------------------------------
def _all_text(doc: Document) -> list[str]:
    """Every text span in the document: body paragraphs + every table cell."""
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return parts


def residual_tags(doc: Document) -> list[str]:
    """Pitfall 1 detection — any leftover Jinja/docxtpl markup or `{{`/`{%` after a
    render means a tag was NOT substituted (a silent non-fill). docxtpl is run-safe by
    design, so this should always be empty for the docxtpl path; the run-replace path
    (Phase 101) is where run-split silent misses actually occur."""
    hits: list[str] = []
    for txt in _all_text(doc):
        if any(tok in txt for tok in ("{{", "}}", "{%", "%}")):
            hits.append(txt.strip()[:80])
    return hits


def assert_integrity(out_path: str, expect_min_rows: int) -> dict:
    """Re-open the produced `.docx` with python-docx (Pitfall 3 gate — raises if the
    file is corrupt / won't open) and report table/row counts + the residual-tag scan.

    `rows` counts EVERY table row across the document (header + grown body rows). The
    caller asserts the exact `header + N` count for the row-growth probe (Pitfall 5).
    """
    doc = Document(out_path)                        # raises if corrupt / won't open
    tables = doc.tables
    rows = sum(len(t.rows) for t in tables)
    residuals = residual_tags(doc)
    return {
        "opened": True,
        "tables": len(tables),
        "rows": rows,
        "rows_ok": rows >= expect_min_rows,
        "residual_tags": residuals,
        "residual_clean": len(residuals) == 0,
    }


# ---------------------------------------------------------------------------
# Corruption-log writer — one pipe-delimited row per named failure mode
# ---------------------------------------------------------------------------
def log_pitfall(name: str, status: str, note: str, *, path: Path | None = None) -> None:
    """Append `<name> | <status> | <note>` to out/corruption.log.

    `name` carries the pitfall label (e.g. "Pitfall 2 | unescaped & < > XML corruption"),
    `status` is one of clean | corrupt | not-exercised | observed, and `note` is the
    where-broke / evidence detail. These rows ARE the Phase 101 UAT seed (SC#4)."""
    path = path or CORRUPTION_LOG
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "a", encoding="utf-8") as fh:
        fh.write(f"{name} | {status} | {note}\n")
