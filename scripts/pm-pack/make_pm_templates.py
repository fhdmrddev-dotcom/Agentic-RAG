"""Phase 104 Plan 01 — reproducible builder for the 2 PM docxtpl templates.

CONTENT/DATA ONLY (PM-01). Builds, via python-docx (no manual Word authoring), the two
``.docx`` templates the headline template-fill workflows render:

  - ``templates/weekly-status-report.docx`` — scalar tags the model emits as cited
    ``FlatScalar``s (``{{ key.value }}`` -> context key ``key``).
  - ``templates/risk-register.docx`` — a 9-column ``{%tr for r in rows %}`` table where
    the model emits 8 cited columns (``{{ r.<col>.value }}``) and the **Score** cell is
    an INLINE Jinja P×I expression over the WORDED probability/impact (NOT a model field,
    NOT backend code).

Mirrors the proven spike-097 conventions (``scripts/spike-097/make_template.py``):
every Jinja tag lives in its OWN single-run paragraph / cell (Pitfall-4-safe — a tag
split across runs raises ``TemplateSyntaxError``), and a ``{%tr %}`` table loop is
authored as a dedicated open-row + content-row + dedicated close-row (docxtpl REPLACES
the whole ``<w:tr>`` row that holds a ``{%tr %}`` tag with the bare ``{% ... %}``
statement, so the loop tags MUST sit in their own rows).

WHY the Score is inline Jinja, not ``{{ r.score }}`` (load-bearing, D-104-4): the
production render context is the SANDBOX-INLINED ``build_context``
(``tool_dispatcher.py:1225-1240``) which applies NO ``numeric_hook`` — so a bare
``{{ r.score }}`` would render BLANK in prod. The inline expression derives the score
deterministically from the cited probability/impact values (Low=1 / Med=2 / High=3,
spike CONCLUSION Condition 3), rendered inside docxtpl's
``SandboxedEnvironment(autoescape=True)`` (``tool_dispatcher.py:1267``) — no backend
change, no model-emitted score field. Confirmed dict access: ``build_context`` rows are
plain dicts ``{col: {"value": ...}}`` so ``r.probability.value`` resolves to
``row["probability"]["value"]`` (Jinja attribute access falls back to item access).

Run from the repo root:
    backend/venv/Scripts/python.exe scripts/pm-pack/make_pm_templates.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

TEMPLATES_DIR = Path(__file__).resolve().parent / "templates"
STATUS_PATH = TEMPLATES_DIR / "weekly-status-report.docx"
RISK_PATH = TEMPLATES_DIR / "risk-register.docx"


# ---------------------------------------------------------------------------
# Weekly Status Report — scalar tags the model emits as cited FlatScalars.
# The TAG NAMES are the contract: each {{ key.value }} -> EmitFieldMap scalar key `key`,
# extracted by parse_docx_template_variables (template_render_service.py:356-413).
# ---------------------------------------------------------------------------
# (heading label, scalar context key) pairs — layout is discretionary, tag names are not.
STATUS_SECTIONS = [
    ("Project", "project_name"),
    ("Reporting Period", "reporting_period"),
    ("Overall Status", "overall_rag_status"),
    ("Summary", "summary"),
    ("Accomplishments This Period", "accomplishments"),
    ("Planned Next Period", "planned_next"),
    ("Risks & Blockers", "risks_blockers"),
    ("Key Milestones & Metrics", "milestones"),
]


def build_status_report() -> Path:
    """Write templates/weekly-status-report.docx (scalar tags under labeled headings)."""
    doc = Document()
    doc.add_heading("Weekly Status Report", level=0)

    for label, key in STATUS_SECTIONS:
        doc.add_heading(label, level=1)
        # Each Jinja tag in its OWN single-run paragraph (Pitfall-4-safe).
        doc.add_paragraph("{{ %s.value }}" % key)

    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(STATUS_PATH))
    # Re-open to confirm integrity.
    Document(str(STATUS_PATH))
    return STATUS_PATH


# ---------------------------------------------------------------------------
# Risk Register — 9-col {%tr%} table; 8 cited columns + inline P×I Score.
# ---------------------------------------------------------------------------
RISK_HEADERS = [
    "Risk ID", "Description", "Category", "Probability", "Impact",
    "Score", "Owner", "Mitigation", "Status",
]

# Worded->numeric P×I score, computed INLINE in Jinja over the SAME cited
# probability/impact cells the model emits. Case- and synonym-tolerant via |lower +
# .get(..., 0); unmapped/blank -> 0 (honest degrade, never a crash). Renders inside
# docxtpl's SandboxedEnvironment (tool_dispatcher.py:1267). This is the ONLY non-.value
# cell — it deliberately is NOT `{{ r.score }}` (which would render blank in prod since
# build_context applies no numeric_hook).
_SCORE_MAP = "{'low':1,'medium':2,'med':2,'high':3}"
SCORE_EXPR = (
    "{{ (%s.get((r.probability.value or '')|trim|lower, 0))"
    " * (%s.get((r.impact.value or '')|trim|lower, 0)) }}" % (_SCORE_MAP, _SCORE_MAP)
)

# Body row tags in the SAME column order as RISK_HEADERS. The model emits 8 cited
# columns; the Score column is the inline P×I expression above.
RISK_BODY_TAGS = [
    "{{ r.id.value }}",
    "{{ r.description.value }}",
    "{{ r.category.value }}",
    "{{ r.probability.value }}",
    "{{ r.impact.value }}",
    SCORE_EXPR,
    "{{ r.owner.value }}",
    "{{ r.mitigation.value }}",
    "{{ r.status.value }}",
]

assert len(RISK_HEADERS) == len(RISK_BODY_TAGS), "header/body column count mismatch"


def build_risk_register() -> Path:
    """Write templates/risk-register.docx (9-col {%tr%} table + inline P×I Score)."""
    doc = Document()
    doc.add_heading("Risk Register", level=0)
    # Scalar context tags — each in its own single-run paragraph.
    doc.add_paragraph("Project: {{ project_name.value }}")
    doc.add_paragraph("Reporting Period: {{ reporting_period.value }}")
    doc.add_paragraph("")  # spacer

    # 4 rows: header + dedicated {%tr for%} row + content row + dedicated {%tr endfor%}
    # row. docxtpl replaces each {%tr %}-bearing row with the bare {% ... %} statement,
    # so the loop tags MUST sit in their own rows (mirror spike make_template.py).
    table = doc.add_table(rows=4, cols=len(RISK_HEADERS))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for i, label in enumerate(RISK_HEADERS):
        header_cells[i].text = label

    # Dedicated loop-open row — tag in the first cell, rest left blank.
    table.rows[1].cells[0].text = "{%tr for r in rows %}"

    # Content row — the row docxtpl repeats once per `r` in `rows`.
    content_cells = table.rows[2].cells
    for i, tag in enumerate(RISK_BODY_TAGS):
        content_cells[i].text = tag

    # Dedicated loop-close row.
    table.rows[3].cells[0].text = "{%tr endfor %}"

    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(RISK_PATH))
    # Re-open to confirm integrity.
    Document(str(RISK_PATH))
    return RISK_PATH


def build() -> list[Path]:
    return [build_status_report(), build_risk_register()]


if __name__ == "__main__":
    made = build()
    print("wrote PM docxtpl templates:")
    for p in made:
        d = Document(str(p))
        print(f"  - {p.name}  ({len(d.tables)} table(s), re-opened OK)")
